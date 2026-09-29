# Core backend: template storage, attendee import, badge rendering, and print exports.
# Start with render_badge() for the full image pipeline and _text_layer() for typography.
from __future__ import annotations

import base64
import csv
import io
import json
import math
import os
import re
import shutil
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

import qrcode
import zxingcpp
from app import google_fonts, typography, security
from app.security import BadgeError, safe_relative_path
from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageCms, ImageColor, ImageOps, UnidentifiedImageError
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as pdf_canvas


# Persistent data lives in .badge_data by default; BADGE_DATA_DIR can point elsewhere.
# Bundled templates are starters; user edits are stored as custom Tier overrides.
PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_ROOT = Path(os.environ.get("BADGE_DATA_DIR", PROJECT_ROOT / ".badge_data"))
BUILTIN_TEMPLATE_ROOT = Path(__file__).resolve().parent / "builtin_templates"
ALLOWED_AVATAR_EXTENSIONS = set(security.IMAGE_EXTENSIONS)


class ExportCancelled(BadgeError):
    """A batch export stopped at a boundary where its output can be discarded."""


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


# Coordinate contract: layout geometry uses millimetres, font sizes use points,
# raster rendering uses pixels at the template PPI, and PDF pages use points.
def mm_to_px(value_mm: float, ppi: int) -> int:
    return max(1, round(float(value_mm) / 25.4 * int(ppi)))


def mm_to_pt(value_mm: float) -> float:
    return float(value_mm) / 25.4 * 72


def event_directory(event_id: str) -> Path:
    if not re.fullmatch(r"[a-zA-Z0-9-]{8,64}", event_id):
        raise BadgeError("Invalid event identifier.")
    path = security.asset_path(DATA_ROOT / "events", event_id)
    if not path.exists():
        raise BadgeError("Event not found.")
    return path


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def manifest_path(root: Path) -> Path:
    path = root / "template.json"
    if not path.is_file():
        raise BadgeError(f"Template is missing template.json: {root.name}")
    return path


# Template boundary: check required settings and referenced assets before accepting a layout.
def _require_asset(root: Path, value: str, extensions: set[str]) -> None:
    path = security.asset_path(root, value)
    if path.suffix.lower() not in extensions or not path.is_file():
        raise BadgeError(f"Missing or unsupported template asset: {value}")


def validate_template(manifest: dict[str, Any], root: Path) -> None:
    if not isinstance(manifest, dict):
        raise BadgeError("Template must be an object.")
    for field in ("id", "name", "version", "profile_picture", "print", "elements"):
        if field not in manifest:
            raise BadgeError(f"Template {root.name} is missing '{field}'.")
    if not isinstance(manifest["profile_picture"], bool):
        raise BadgeError("Template profile_picture must be true or false.")
    for key in ("id", "name"):
        if not isinstance(manifest[key], str) or not 1 <= len(manifest[key]) <= 120:
            raise BadgeError(f"Template {key} must be text with at most 120 characters.")
    if "builtin_source" in manifest and (not isinstance(manifest["builtin_source"], str)
                                          or not 1 <= len(manifest["builtin_source"]) <= 120):
        raise BadgeError("Template builtin_source must be text with at most 120 characters.")
    if "priority" in manifest:
        security.finite_number(manifest["priority"], -1000, 1000, "template priority")
    printing = manifest["print"]
    if not isinstance(printing, dict):
        raise BadgeError("Template print settings must be an object.")
    for key, low, high in (("width_mm", .001, 2000), ("height_mm", .001, 2000),
                           ("bleed_mm", 0, 100), ("safe_mm", 0, 1000), ("ppi", 72, 1200)):
        security.finite_number(printing.get(key), low, high, f"print.{key}")
    ppi = int(printing["ppi"])
    bleed = mm_to_px(printing["bleed_mm"], ppi)
    width = mm_to_px(printing["width_mm"], ppi) + 2 * bleed
    height = mm_to_px(printing["height_mm"], ppi) + 2 * bleed
    security.check_dimensions(width, height)
    elements = manifest["elements"]
    if not isinstance(elements, list) or len(elements) > security.MAX_LAYERS:
        raise BadgeError("Template elements must be a list with at most 100 layers.")
    background = manifest.get("background", {})
    if not isinstance(background, dict):
        raise BadgeError("Template background must be an object.")
    if background.get("asset"):
        _require_asset(root, background["asset"], ALLOWED_AVATAR_EXTENSIONS)
    total_pixels = width * height
    for element in elements:
        if not isinstance(element, dict) or element.get("type") not in ("text", "masked-image", "qr", "artwork"):
            raise BadgeError("Template has an unsupported element type.")
        for key in ("x_mm", "y_mm", "width_mm", "height_mm"):
            if key not in element:
                raise BadgeError(f"Template element is missing {key}.")
        total_pixels += security.validate_element_geometry(element, ppi)
        if total_pixels > security.MAX_LAYER_PIXELS:
            raise BadgeError("The combined area of the template layers is too large.")
        if element["type"] == "qr" and element.get("symbology", "qr") not in ("qr", "data-matrix", "aztec"):
            raise BadgeError("Unsupported barcode type. Choose QR Code, Data Matrix, or Aztec.")
        if "name" in element and (not isinstance(element["name"], str) or len(element["name"]) > 120):
            raise BadgeError("Layer names must be text with at most 120 characters.")
        if element["type"] == "text" and "text" in element:
            if not isinstance(element["text"], str) or len(element["text"]) > 8192:
                raise BadgeError("Custom text must contain at most 8192 characters.")
        elif element["type"] in ("text", "masked-image", "qr"):
            if not isinstance(element.get("field"), str) or not element["field"]:
                raise BadgeError("Template element must have a field binding.")
        if element["type"] == "text" and element.get("font_language", "auto") not in ("auto", "zh-Hans", "zh-Hant", "ja", "ko"):
            raise BadgeError("Choose a supported font language.")
        if element["type"] == "artwork":
            _require_asset(root, element.get("asset"), ALLOWED_AVATAR_EXTENSIONS)
        for key in ("visible", "lock_aspect_ratio", "snap_to_edges"):
            if key in element and not isinstance(element[key], bool):
                raise BadgeError(f"Invalid {key} setting.")
        if "aspect_ratio" in element:
            security.finite_number(element["aspect_ratio"], 1 / security.MAX_SIDE, security.MAX_SIDE, "aspect ratio")
        if element.get("mask_asset"):
            _require_asset(root, element["mask_asset"], ALLOWED_AVATAR_EXTENSIONS)
        if element.get("font_asset"):
            _require_asset(root, element["font_asset"], {".ttf", ".otf"})
    has_avatar = any(element["type"] == "masked-image" for element in elements)
    if manifest["profile_picture"] != has_avatar:
        raise BadgeError("Template profile_picture does not match its avatar element.")
    fonts = manifest.get("fonts", [])
    if not isinstance(fonts, list) or len(fonts) > 100:
        raise BadgeError("Template fonts must be a list with at most 100 families.")
    seen_ids = set()
    for font in fonts:
        if not isinstance(font, dict) or not isinstance(font.get("name"), str) or not 1 <= len(font["name"]) <= 120:
            raise BadgeError("Template has an invalid font family.")
        security.validate_font_id(font.get("id"))
        if font["id"] in seen_ids:
            raise BadgeError("Template has duplicate font IDs.")
        seen_ids.add(font["id"])
        for style in ("regular", "bold", "italic", "bold_italic"):
            if font.get(style) is not None:
                _require_asset(root, font[style], {".ttf", ".otf"})


# Load bundled and custom manifests. This also migrates older saved formats on disk,
# so calling it may rename legacy assets or rewrite legacy manifests.
def discover_templates() -> dict[str, dict[str, Any]]:
    roots = [BUILTIN_TEMPLATE_ROOT, DATA_ROOT / "templates"]
    templates: dict[str, dict[str, Any]] = {}
    for base in roots:
        if not base.exists():
            continue
        for candidate in base.rglob("template.json"):
            root = candidate.parent
            manifest = read_json(candidate)
            # Upgrade locally saved templates produced before the simplified one-Tier schema.
            if "profile_picture" not in manifest:
                legacy_tiers = manifest.pop("supported_tiers", [])
                if legacy_tiers:
                    manifest["name"] = legacy_tiers[0]
                manifest["profile_picture"] = any(element.get("type") == "masked-image" for element in manifest.get("elements", []))
                write_json(candidate, manifest)
            if manifest.get("id", "").startswith("tier-") and str(manifest.get("name", "")).endswith(" Template"):
                manifest["name"] = str(manifest["name"])[:-9]
                write_json(candidate, manifest)
            legacy_background = manifest.get("background", {}).get("asset")
            if legacy_background and Path(legacy_background).name in {"background.jpg", "background.jpeg", "background.png"}:
                old_background = security.asset_path(root, legacy_background)
                renamed = background_filename(manifest["name"], old_background.suffix)
                new_background = root / renamed
                if old_background.is_file() and not new_background.exists():
                    old_background.rename(new_background)
                    manifest["background"]["asset"] = renamed
                    write_json(candidate, manifest)
            validate_template(manifest, root)
            identifier = manifest["id"]
            if identifier in templates:
                raise BadgeError(f"Duplicate template ID: {identifier}")
            templates[identifier] = {"manifest": manifest, "root": root}
    return templates


def available_templates(templates: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """A custom copy replaces its starter even when its display name changes.

    Keep the original in discover_templates() for existing records referencing its
    ID, but omit it from the Tier list, new imports, and format exports.
    """
    replaced = {item["manifest"].get("builtin_source") for item in templates.values()
                if item["root"].is_relative_to(DATA_ROOT / "templates")}
    return {identifier: item for identifier, item in templates.items()
            if not (item["root"].is_relative_to(BUILTIN_TEMPLATE_ROOT) and identifier in replaced)}


def editable_template_manifest(template: dict[str, Any]) -> dict[str, Any]:
    manifest = json.loads(json.dumps(template["manifest"]))
    if template["root"].is_relative_to(BUILTIN_TEMPLATE_ROOT):
        manifest["builtin_source"] = manifest["id"]
    elif "builtin_source" not in manifest:
        # Older custom overrides kept the starter's name but had no source ID.
        for path in BUILTIN_TEMPLATE_ROOT.rglob("template.json"):
            starter = read_json(path)
            if starter["name"].casefold() == manifest["name"].casefold():
                manifest["builtin_source"] = starter["id"]
                break
    return manifest


def template_slug(tier: str) -> str:
    cleaned = re.sub(r"[^a-z0-9]+", "-", tier.casefold()).strip("-")
    if not cleaned:
        raise BadgeError("Tier must include at least one letter or number.")
    return cleaned[:60]


def background_filename(tier: str, extension: str) -> str:
    stem = template_slug(tier)
    if stem == "attendees":
        stem = "attendee"
    return f"{stem}-bg{extension.lower()}"


@security.image_task
def create_tier_template(tier: str, image_filename: str, image_bytes: security.ImageSource, avatar_enabled: bool) -> dict[str, Any]:
    """Create a user-facing, one-image-per-tier template with editable default slots."""
    tier = tier.strip()
    if not tier or len(tier) > 120:
        raise BadgeError("Tier is required and must be 120 characters or fewer.")
    extension = Path(image_filename).suffix.lower()
    if extension not in ALLOWED_AVATAR_EXTENSIONS:
        raise BadgeError("Template artwork must be a JPG, JPEG, or PNG image.")


    template_id = f"tier-{template_slug(tier)}"
    template_root = DATA_ROOT / "templates" / template_id
    if template_root.exists():
        raise BadgeError("A custom Tier with this name already exists. Use Edit Tier or upload a new background instead.")
    staging = DATA_ROOT / f".staging-{uuid.uuid4()}"
    staging.mkdir(parents=True, exist_ok=False)
    try:
        background_name = background_filename(tier, extension)
        width_px, height_px, ppi = security.normalize_image_file(
            image_filename, image_bytes, staging / background_name, background=True)
        width_mm = round(width_px / ppi * 25.4, 2)
        height_mm = round(height_px / ppi * 25.4, 2)
        if width_mm <= 0 or height_mm <= 0:
            raise BadgeError("Template artwork has invalid dimensions.")
        name_width = width_mm * .8
        elements: list[dict[str, Any]] = [
            {"type": "text", "field": "tier", "x_mm": round(width_mm * .1, 2), "y_mm": round(height_mm * .04, 2), "width_mm": round(name_width, 2), "height_mm": round(height_mm * .06, 2), "font_size_pt": 12, "min_font_size_pt": 8, "align": "center", "color": "#FFFFFF"},
            {"type": "text", "field": "display_name", "x_mm": round(width_mm * .1, 2), "y_mm": round(height_mm * .67, 2), "width_mm": round(name_width, 2), "height_mm": round(height_mm * .12, 2), "font_size_pt": 24, "min_font_size_pt": 11, "align": "center", "color": "#FFFFFF"},
            {"type": "text", "field": "ticket_id", "x_mm": round(width_mm * .1, 2), "y_mm": round(height_mm * .81, 2), "width_mm": round(name_width, 2), "height_mm": round(height_mm * .06, 2), "font_size_pt": 12, "min_font_size_pt": 8, "align": "center", "color": "#FFFFFF"},
            {"type": "qr", "field": "qr_token", "x_mm": round(width_mm * .75, 2), "y_mm": round(height_mm * .86, 2), "width_mm": round(width_mm * .16, 2), "height_mm": round(width_mm * .16, 2), "quiet_zone_modules": 4},
        ]
        if avatar_enabled:
            avatar_size = min(width_mm * .5, height_mm * .42)
            elements.insert(1, {"type": "masked-image", "field": "avatar_asset", "x_mm": round((width_mm - avatar_size) / 2, 2), "y_mm": round(height_mm * .15, 2), "width_mm": round(avatar_size, 2), "height_mm": round(avatar_size, 2), "fit": "cover", "missing": "transparent"})
        manifest = {
            "id": template_id,
            "name": tier,
            "version": now(),
            "priority": 100,
            "profile_picture": avatar_enabled,
            "print": {"width_mm": width_mm, "height_mm": height_mm, "bleed_mm": 0, "safe_mm": 0, "ppi": ppi},
            "background": {"asset": background_name},
            "elements": elements,
        }
        write_json(staging / "template.json", manifest)
        validate_template(manifest, staging)
        template_root.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(staging), str(template_root))
        return {"id": template_id, "name": manifest["name"], "tier": tier, "width_mm": width_mm, "height_mm": height_mm, "ppi": ppi}
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise


# Upload background: preserve the existing layer layout by scaling its millimetre coordinates
# and font sizes to the new artwork dimensions.
@security.image_task
def update_template_background(template_id: str, image_filename: str, image_bytes: security.ImageSource) -> dict[str, Any]:
    """Replace the artwork for one active Tier while retaining its layout manifest."""
    extension = Path(image_filename).suffix.lower()
    if extension not in ALLOWED_AVATAR_EXTENSIONS:
        raise BadgeError("Background artwork must be a JPG, JPEG, or PNG image.")
    templates = discover_templates()
    source = templates.get(template_id)
    if not source:
        raise BadgeError("Template not found.")
    manifest = editable_template_manifest(source)
    tier = manifest["name"]
    target_id = f"tier-{template_slug(tier)}"
    target_root = DATA_ROOT / "templates" / target_id
    staging = DATA_ROOT / f".staging-{uuid.uuid4()}"
    try:
        shutil.copytree(source["root"], staging)
        background_name = background_filename(tier, extension)
        width_px, height_px, ppi = security.normalize_image_file(
            image_filename, image_bytes, staging / background_name, background=True)
        width_mm = round(width_px / ppi * 25.4, 2)
        height_mm = round(height_px / ppi * 25.4, 2)
        if width_mm <= 0 or height_mm <= 0:
            raise BadgeError("Template artwork has invalid dimensions.")
        old_print = manifest["print"]
        scale_x = width_mm / float(old_print["width_mm"])
        scale_y = height_mm / float(old_print["height_mm"])
        font_scale = (scale_x + scale_y) / 2
        for element in manifest["elements"]:
            element["x_mm"] = round(float(element["x_mm"]) * scale_x, 2)
            element["width_mm"] = round(float(element["width_mm"]) * scale_x, 2)
            element["y_mm"] = round(float(element["y_mm"]) * scale_y, 2)
            element["height_mm"] = round(float(element["height_mm"]) * scale_y, 2)
            for font_field in ("font_size_pt", "min_font_size_pt"):
                if font_field in element:
                    element[font_field] = round(float(element[font_field]) * font_scale, 2)
        manifest.update({
            "id": target_id,
            "name": tier,
            "version": now(),
            "priority": 100,
            "print": {"width_mm": width_mm, "height_mm": height_mm, "bleed_mm": 0, "safe_mm": 0, "ppi": ppi},
            "background": {"asset": background_name},
        })
        write_json(staging / "template.json", manifest)
        validate_template(manifest, staging)
        if target_root.exists():
            shutil.rmtree(target_root)
        target_root.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(staging), str(target_root))
        return {"id": target_id, "name": manifest["name"], "tier": tier}
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise


# Edit layout bootstrap data: include only same-Tier attendees with an existing avatar file.
def editor_template(template_id: str, event_id: str | None = None) -> dict[str, Any]:
    template = discover_templates().get(template_id)
    if not template:
        raise BadgeError("Template not found.")
    manifest = template["manifest"]
    attendees: list[dict[str, str]] = []
    if event_id:
        event_dir, event = load_event(event_id)
        for record in event.get("records", []):
            reference = str(record.get("avatar_asset", ""))
            if str(record.get("tier", "")).casefold() != str(manifest["name"]).casefold() or not reference:
                continue
            try:
                if avatar_path(event_dir, reference).is_file():
                    attendees.append({"ticket_id": record["ticket_id"], "display_name": record["display_name"], "avatar_url": f"/api/events/{event_id}/records/{record['ticket_id']}/avatar"})
            except BadgeError:
                continue
    return {"id": template_id, "manifest": manifest, "preview_attendees": attendees}


def event_avatar_preview(event_id: str, ticket_id: str) -> Path:
    event_dir, event = load_event(event_id)
    record = next((item for item in event["records"] if item["ticket_id"] == ticket_id), None)
    if not record:
        raise BadgeError("Ticket ID not found.")
    reference = str(record.get("avatar_asset", ""))
    if not reference:
        raise BadgeError("This attendee has no avatar.")
    path = avatar_path(event_dir, reference)
    if not path.is_file():
        raise BadgeError("This attendee avatar is missing.")
    return path


# Shared barcode encoder for editor previews and finished badges.
# Keep module generation here so QR, Data Matrix, and Aztec behave consistently.
def barcode_image(value: str, symbology: str = "qr", foreground: str = "#000000", background: str = "#ffffff", quiet_zone_modules: int = 4, error_correction: str = "M") -> Image.Image:
    """Return a barcode at one pixel per module, including its quiet zone."""
    if symbology not in {"qr", "data-matrix", "aztec"}:
        raise BadgeError("Unsupported barcode type. Choose QR Code, Data Matrix, or Aztec.")
    if not value:
        raise BadgeError("Barcode content cannot be empty.")
    if len(value) > 8192:
        raise BadgeError("Barcode content is too long.")
    try:
        ImageColor.getrgb(foreground)
        ImageColor.getrgb(background)
        border = max(4, min(16, int(quiet_zone_modules)))
        if symbology == "qr":
            corrections = {"L": qrcode.constants.ERROR_CORRECT_L, "M": qrcode.constants.ERROR_CORRECT_M, "Q": qrcode.constants.ERROR_CORRECT_Q, "H": qrcode.constants.ERROR_CORRECT_H}
            code = qrcode.QRCode(error_correction=corrections.get(error_correction.upper(), qrcode.constants.ERROR_CORRECT_M), border=border, box_size=1)
            code.add_data(value)
            code.make(fit=True)
            return code.make_image(fill_color=foreground, back_color=background).convert("RGB")
        formats = {"data-matrix": zxingcpp.BarcodeFormat.DataMatrix, "aztec": zxingcpp.BarcodeFormat.Aztec}
        code = zxingcpp.create_barcode(value, formats[symbology])
        raster = code.to_image(scale=1, add_quiet_zones=False)
        height, width = memoryview(raster).shape
        modules = Image.frombytes("L", (width, height), bytes(raster))
        modules = ImageOps.expand(modules, border=border, fill=255)
        return ImageOps.colorize(modules, black=foreground, white=background)
    except (ValueError, TypeError, OverflowError, RuntimeError, qrcode.exceptions.DataOverflowError) as exc:
        raise BadgeError("Cannot generate barcode. Check the content length, colors, and barcode settings.") from exc


def qr_preview(value: str, foreground: str = "#000000", background: str = "#ffffff", quiet_zone_modules: int = 4, symbology: str = "qr", error_correction: str = "M") -> bytes:
    image = barcode_image(value, symbology, foreground, background, quiet_zone_modules, error_correction)
    image = image.resize((image.width * 8, image.height * 8), Image.Resampling.NEAREST)
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


def template_background_path(template_id: str) -> Path | None:
    template = discover_templates().get(template_id)
    if not template:
        raise BadgeError("Template not found.")
    asset = template["manifest"].get("background", {}).get("asset")
    if not asset:
        return None
    path = security.asset_path(template["root"], asset)
    if not path.is_file():
        raise BadgeError("Template background asset is missing.")
    return path


def template_asset_path(template_id: str, asset_path: str) -> Path:
    if safe_relative_path(asset_path).suffix.lower() not in security.ASSET_EXTENSIONS:
        raise BadgeError("Unsupported template asset type.")
    template = discover_templates().get(template_id)
    if not template:
        raise BadgeError("Template not found.")
    path = security.asset_path(template["root"], asset_path)
    if not path.is_file():
        raise BadgeError("Template asset is missing.")
    return path


# Save layout core: copy assets to a staging folder, validate the new manifest,
# then replace the custom Tier folder. The bundled starter remains available.
def save_template_layout(template_id: str, elements: list[dict[str, Any]], fonts: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    template = discover_templates().get(template_id)
    if not template:
        raise BadgeError("Template not found.")
    manifest = editable_template_manifest(template)
    if not isinstance(elements, list):
        raise BadgeError("Template elements must be a list.")
    manifest["elements"] = elements
    if fonts is not None:
        manifest["fonts"] = fonts
    tier = manifest["name"]
    target_id = f"tier-{template_slug(tier)}"
    target_root = DATA_ROOT / "templates" / target_id
    staging = DATA_ROOT / f".staging-{uuid.uuid4()}"
    try:
        shutil.copytree(template["root"], staging)
        manifest.update({"id": target_id, "name": tier, "version": now(), "priority": 100})
        write_json(staging / "template.json", manifest)
        validate_template(manifest, staging)
        if target_root.exists():
            shutil.rmtree(target_root)
        target_root.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(staging), str(target_root))
        return {"id": target_id, "tier": tier}
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise


@security.image_task
def add_template_asset(template_id: str, kind: str, filename: str, content: bytes) -> dict[str, Any]:
    """Store an editor asset in a custom Template override and return its safe path."""
    directories = {"artwork": "artwork", "mask": "masks"}
    if kind not in directories:
        raise BadgeError("Unsupported template asset type.")
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_AVATAR_EXTENSIONS:
        raise BadgeError("Artwork and masks must be JPG, JPEG, or PNG images.")
    content, width_px, height_px, _ = security.normalize_image(filename, content)
    template = discover_templates().get(template_id)
    if not template:
        raise BadgeError("Template not found.")
    manifest = editable_template_manifest(template)
    target_id = f"tier-{template_slug(manifest['name'])}"
    target_root = DATA_ROOT / "templates" / target_id
    staging = DATA_ROOT / f".staging-{uuid.uuid4()}"
    asset_name = f"{directories[kind]}/{uuid.uuid4().hex}{extension}"
    try:
        shutil.copytree(template["root"], staging)
        destination = security.asset_path(staging, asset_name)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(content)
        manifest.update({"id": target_id, "version": now(), "priority": 100})
        write_json(staging / "template.json", manifest)
        validate_template(manifest, staging)
        if target_root.exists():
            shutil.rmtree(target_root)
        target_root.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(staging), str(target_root))
        return {"id": target_id, "asset": asset_name, "width_px": width_px, "height_px": height_px}
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise


# Uploaded and Google fonts share asset storage. Google downloads retain their OFL license.
# It stores the font asset immediately; assigning it to a text layer is saved separately.
def add_template_font(template_id: str, family_name: str, style: str, filename: str, content: bytes, *, source_info: dict | None = None, license_content: bytes | None = None) -> dict[str, Any]:
    family_name = family_name.strip()
    if not family_name or len(family_name) > 120:
        raise BadgeError("Font family name is required.")
    if style not in {"regular", "bold", "italic", "bold_italic"}:
        raise BadgeError("Choose a valid font style.")
    extension = Path(filename).suffix.lower()
    if extension not in {".ttf", ".otf"}:
        raise BadgeError("Fonts must be TTF or OTF files.")
    security.validate_font(filename, content)
    template = discover_templates().get(template_id)
    if not template:
        raise BadgeError("Template not found.")
    manifest = editable_template_manifest(template)
    fonts = manifest.setdefault("fonts", [])
    family = next((item for item in fonts if str(item.get("name", "")).casefold() == family_name.casefold()), None)
    if family is None:
        family = {"id": f"font-{uuid.uuid4().hex[:12]}", "name": family_name}
        fonts.append(family)
    target_id = f"tier-{template_slug(manifest['name'])}"
    target_root = DATA_ROOT / "templates" / target_id
    staging = DATA_ROOT / f".staging-{uuid.uuid4()}"
    security.validate_font_id(family["id"])
    asset_name = f"fonts/{family['id']}-{style}{extension}"
    try:
        shutil.copytree(template["root"], staging)
        destination = security.asset_path(staging, asset_name)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(content)
        family[style] = asset_name
        # Track each file separately: uploading another style does not inherit Google's license.
        origins = family.setdefault("origins", {})
        origins[style] = source_info or {"source": "upload"}
        if license_content:
            license_asset = f"fonts/{family['id']}-{style}-OFL.txt"
            security.asset_path(staging, license_asset).write_bytes(security.normalize_asset(license_asset, license_content))
            origins[style]["license_asset"] = license_asset
        manifest.update({"id": target_id, "version": now(), "priority": 100})
        write_json(staging / "template.json", manifest)
        validate_template(manifest, staging)
        if target_root.exists():
            shutil.rmtree(target_root)
        target_root.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(staging), str(target_root))
        return {"id": target_id, "font": family, "fonts": fonts}
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise


def add_google_font(template_id: str, value: str, style: str = "regular") -> dict[str, Any]:
    if style not in {"regular", "bold", "italic", "bold_italic"}:
        raise BadgeError("Choose a valid font style.")
    try:
        info = google_fonts.ensure_font(value, DATA_ROOT / "google-fonts", italic="italic" in style, bold="bold" in style)
        return add_template_font(template_id, info["name"], style, "google.ttf", info["path"].read_bytes(),
                                 source_info={k: v for k, v in info.items() if k not in {"path", "license_path"}},
                                 license_content=info["license_path"].read_bytes())
    except google_fonts.FontError as exc:
        raise BadgeError(str(exc)) from exc


def update_template_settings(template_id: str, name: str, profile_picture: bool) -> dict[str, Any]:
    installed = discover_templates()
    template = installed.get(template_id)
    name = name.strip()
    if not template or not name:
        raise BadgeError("Template and Tier name are required.")
    manifest = editable_template_manifest(template)
    if name.casefold() != manifest["name"].casefold() and any(
            item["manifest"]["name"].casefold() == name.casefold()
            for item in available_templates(installed).values()):
        raise BadgeError("Another Tier already uses this name. Choose a different name.")
    elements = manifest["elements"]
    avatars = [element for element in elements if element.get("type") == "masked-image"]
    if profile_picture and not avatars:
        printing = manifest["print"]
        size = min(float(printing["width_mm"]) * .5, float(printing["height_mm"]) * .42)
        elements.insert(0, {"type": "masked-image", "field": "avatar_asset", "x_mm": round((float(printing["width_mm"]) - size) / 2, 2), "y_mm": round(float(printing["height_mm"]) * .15, 2), "width_mm": round(size, 2), "height_mm": round(size, 2), "fit": "cover", "missing": "transparent"})
    if not profile_picture:
        elements[:] = [element for element in elements if element.get("type") != "masked-image"]
    target_id = f"tier-{template_slug(name)}"
    target_root = DATA_ROOT / "templates" / target_id
    if target_root.exists() and template["root"] != target_root:
        raise BadgeError("Another custom Tier already uses this name. Choose a different name.")
    staging = DATA_ROOT / f".staging-{uuid.uuid4()}"
    try:
        shutil.copytree(template["root"], staging)
        manifest.update({"id": target_id, "name": name, "profile_picture": profile_picture, "version": now(), "priority": 100})
        write_json(staging / "template.json", manifest); validate_template(manifest, staging)
        source_root = template["root"]
        if target_root.exists(): shutil.rmtree(target_root)
        if source_root.is_relative_to(DATA_ROOT / "templates") and source_root != target_root: shutil.rmtree(source_root)
        target_root.parent.mkdir(parents=True, exist_ok=True); shutil.move(str(staging), str(target_root))
        return {"id": target_id, "name": name, "profile_picture": profile_picture}
    except Exception:
        shutil.rmtree(staging, ignore_errors=True); raise


def delete_custom_template(template_id: str) -> None:
    template = discover_templates().get(template_id)
    if not template:
        raise BadgeError("Template not found.")
    root = template["root"]
    if not root.is_relative_to(DATA_ROOT / "templates"):
        raise BadgeError("Built-in Tier templates cannot be deleted.")
    shutil.rmtree(root)


# Download all formats JSON: package active manifests plus their assets for another computer.
@security.image_task
def template_set_json() -> bytes:
    """Export every active Tier as one self-contained JSON document.

    Each Tier carries its own manifest and base64-encoded assets, so the document can
    be restored without referring to files on the original computer.
    """
    installed = available_templates(discover_templates())
    active: list[dict[str, Any]] = []
    total_bytes = 0
    total_assets = 0
    for tier in sorted({item["manifest"]["name"] for item in installed.values()}, key=str.casefold):
        item = template_for_tier(tier, installed)
        if not item:
            raise BadgeError(f"Tier '{tier}' has more than one active template.")
        assets = []
        for path in sorted(item["root"].rglob("*")):
            if not path.is_file() or path.name == "template.json":
                continue
            relative = path.relative_to(item["root"]).as_posix()
            safe_path = security.asset_path(item["root"], relative)
            if safe_path.suffix.lower() not in security.ASSET_EXTENSIONS:
                continue
            if safe_path.stat().st_size > security.MAX_TEMPLATE_BYTES:
                raise BadgeError("Template asset is too large to export.")
            content = security.normalize_asset(relative, safe_path.read_bytes())
            total_bytes += len(content)
            total_assets += 1
            if total_bytes > security.MAX_TEMPLATE_BYTES or total_assets > security.MAX_ASSETS:
                raise BadgeError("Template set exceeds the asset export limit.")
            assets.append({"path": relative, "base64_data": base64.b64encode(content).decode("ascii")})
        active.append({"manifest": item["manifest"], "assets": assets})
        if len(active) > 100:
            raise BadgeError("Export at most 100 Tiers in one template set.")
    package = {"format_version": 1, "exported_at": now(), "templates": active}
    payload = json.dumps(package, ensure_ascii=False, indent=2).encode("utf-8")
    security.check_size(payload, security.MAX_TEMPLATE_BYTES, "Template JSON")
    return payload


@security.image_task
def import_template_set_json(payload: bytes) -> dict[str, Any]:
    """Restore a self-contained JSON template set without deleting unrelated Tiers."""
    security.check_size(payload, security.MAX_TEMPLATE_BYTES, "Template JSON")
    try:
        package = json.loads(payload.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError, RecursionError) as exc:
        raise BadgeError("Template format must be a valid UTF-8 JSON file.") from exc
    if not isinstance(package, dict) or package.get("format_version") != 1:
        raise BadgeError("This is not a supported template format JSON file.")
    entries = package.get("templates")
    if not isinstance(entries, list) or not entries:
        raise BadgeError("Template format JSON does not contain any Tiers.")
    if len(entries) > 100:
        raise BadgeError("A template set may contain at most 100 Tiers.")
    total_bytes = 0
    total_assets = 0
    staging = DATA_ROOT / f".staging-template-set-{uuid.uuid4()}"
    target_root = DATA_ROOT / "templates"
    prepared: list[tuple[Path, Path, str]] = []
    seen_names: set[str] = set()
    seen_ids: set[str] = set()
    try:
        staging.mkdir(parents=True, exist_ok=False)
        for entry in entries:
            if not isinstance(entry, dict) or not isinstance(entry.get("manifest"), dict) or not isinstance(entry.get("assets"), list):
                raise BadgeError("Every Tier in the template format must include a manifest and assets.")
            manifest = json.loads(json.dumps(entry["manifest"]))
            name = str(manifest.get("name", "")).strip()
            if not name:
                raise BadgeError("A template format contains a Tier without a name.")
            name_key = name.casefold()
            if name_key in seen_names:
                raise BadgeError(f"The template format contains Tier '{name}' more than once.")
            seen_names.add(name_key)
            target_id = f"tier-{template_slug(name)}"
            if target_id in seen_ids:
                raise BadgeError("Tier names must produce distinct folder names.")
            seen_ids.add(target_id)
            root = staging / target_id
            root.mkdir(parents=True, exist_ok=False)
            written_assets: set[str] = set()
            total_assets += len(entry["assets"])
            if total_assets > security.MAX_ASSETS:
                raise BadgeError("Too many assets in the template set.")
            for asset in entry["assets"]:
                if not isinstance(asset, dict) or not isinstance(asset.get("path"), str) or not isinstance(asset.get("base64_data"), str):
                    raise BadgeError(f"Tier '{name}' has an invalid asset.")
                relative = safe_relative_path(asset["path"])
                relative_text = relative.as_posix()
                if relative.name.lower() == "template.json" or relative_text.casefold() in written_assets:
                    raise BadgeError(f"Tier '{name}' has a duplicate or reserved asset path.")
                try:
                    content = base64.b64decode(asset["base64_data"], validate=True)
                except ValueError as exc:
                    raise BadgeError(f"Tier '{name}' has an invalid encoded asset.") from exc
                content = security.normalize_asset(relative_text, content)
                total_bytes += len(content)
                if total_bytes > security.MAX_TEMPLATE_BYTES:
                    raise BadgeError("Decoded template assets exceed the 128 MiB limit.")
                destination = security.asset_path(root, relative_text)
                destination.parent.mkdir(parents=True, exist_ok=True)
                destination.write_bytes(content)
                written_assets.add(relative_text.casefold())
            manifest.update({"id": target_id, "name": name, "version": now(), "priority": 100})
            write_json(root / "template.json", manifest)
            validate_template(manifest, root)
            prepared.append((root, target_root / target_id, name))
        for root, final, _ in prepared:
            if final.exists():
                shutil.rmtree(final)
            final.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(root), str(final))
        return {"imported": len(prepared), "tiers": [name for _, _, name in prepared]}
    finally:
        shutil.rmtree(staging, ignore_errors=True)


# Resolve Tier names case-insensitively. A custom override wins by priority;
# equal-priority matches are ambiguous and return None.
def template_for_tier(tier: str, templates: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
    tier_key = tier.casefold()
    matches = [item for item in available_templates(templates).values()
               if str(item["manifest"]["name"]).casefold() == tier_key]
    if not matches:
        return None
    highest = max(item["manifest"].get("priority", 0) for item in matches)
    selected = [item for item in matches if item["manifest"].get("priority", 0) == highest]
    return selected[0] if len(selected) == 1 else None


def avatar_index(assets_directory: Path) -> dict[str, str]:
    """Map ticket-ID-shaped filenames to their relative asset paths.

    CSV does not carry image paths. A single `S-002.png`, `S-002.jpg`, or `S-002.jpeg`
    anywhere inside the chosen avatar folder is associated with ticket ID `S-002`.
    """
    matches: dict[str, list[Path]] = {}
    for candidate in assets_directory.rglob("*"):
        if candidate.is_file() and candidate.suffix.lower() in ALLOWED_AVATAR_EXTENSIONS:
            matches.setdefault(candidate.stem.casefold(), []).append(candidate.relative_to(assets_directory))
    indexed: dict[str, str] = {}
    for ticket_key, paths in matches.items():
        if len(paths) > 1:
            names = ", ".join(path.as_posix() for path in paths)
            raise BadgeError(f"More than one avatar matches '{ticket_key}': {names}. Keep only one.")
        indexed[ticket_key] = paths[0].as_posix()
    return indexed


def avatar_path(event_dir: Path, avatar_reference: str) -> Path:
    return security.asset_path(event_dir / "assets", avatar_reference)


def prepare_record(record: dict[str, Any], row_number: int, templates: dict[str, dict[str, Any]], avatars: dict[str, str]) -> None:
    for field in ("ticket_id", "display_name", "tier", "qr_token"):
        if not (str(record.get(field, "")) if field == "qr_token" else str(record.get(field, "")).strip()):
            raise BadgeError(f"CSV row {row_number}: {field} is required.")
    if any(len(str(value)) > 8192 for value in record.values()):
        raise BadgeError(f"CSV row {row_number}: a field exceeds 8192 characters.")
    template = template_for_tier(record["tier"], templates)
    if not template:
        raise BadgeError(f"CSV row {row_number}: no single template matches tier '{record['tier']}'.")
    record["template_id"] = template["manifest"]["id"]
    record["avatar_asset"] = avatars.get(record["ticket_id"].casefold(), "")


# XLSX is converted to CSV first so both file types share validation and avatar matching.
def attendee_file_to_csv(filename: str, content: bytes) -> bytes:
    """Convert the first XLSX worksheet to UTF-8 CSV for the shared importer."""
    security.check_size(content, security.MAX_SPREADSHEET_BYTES, "Attendee file")
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        return content
    if extension != ".xlsx":
        raise BadgeError("Choose a CSV or XLSX file.")
    workbook = None
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or sum(item.file_size for item in members) > 64 * security.MIB:
                raise BadgeError("The XLSX workbook expands beyond the import limit.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if not workbook.worksheets:
            raise BadgeError("The XLSX file must contain a worksheet.")
        worksheet = workbook.worksheets[0]
        if (worksheet.max_row or 0) > security.MAX_RECORDS + 1 or (worksheet.max_column or 0) > 64:
            raise BadgeError("The XLSX worksheet is too large; use at most 10,000 rows and 64 columns.")
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        for index, row in enumerate(worksheet.iter_rows(), start=1):
            if index > security.MAX_RECORDS + 1 or len(row) > 64:
                raise BadgeError("The XLSX worksheet is too large.")
            values = []
            for cell in row:
                if cell.data_type in {"f", "e"}:
                    raise BadgeError(f"XLSX cell {cell.coordinate}: replace formulas or spreadsheet errors with plain values before importing.")
                value = cell.value
                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                values.append("" if value is None else str(value))
            writer.writerow(values)
            if output.tell() > security.MAX_SPREADSHEET_BYTES:
                raise BadgeError("The converted spreadsheet is too large.")
        return output.getvalue().encode("utf-8")
    except BadgeError:
        raise
    except Exception as exc:
        raise BadgeError("The XLSX file could not be read. Upload a valid, unencrypted .xlsx workbook.") from exc
    finally:
        if workbook is not None:
            workbook.close()


# Import core: parse rows, match avatars by ticket ID, select templates, and write event.json.
# A failed import removes its partially created event folder.
@security.image_task
def create_event(csv_bytes: bytes, avatar_files: list[tuple[str, security.ImageSource]]) -> dict[str, Any]:
    security.check_size(csv_bytes, security.MAX_SPREADSHEET_BYTES, "Attendee CSV")
    if len(avatar_files) > security.MAX_ASSETS or sum(security.source_size(blob) for _, blob in avatar_files) > security.MAX_REQUEST_BYTES:
        raise BadgeError("The avatar import is too large. Use fewer or smaller images.")
    try:
        decoded = csv_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise BadgeError("CSV must be UTF-8 encoded.") from exc
    reader = csv.DictReader(io.StringIO(decoded))
    required = {"ticket_id", "display_name", "tier", "qr_token"}
    if not reader.fieldnames or not required.issubset(set(reader.fieldnames)):
        raise BadgeError("CSV must include ticket_id, display_name, tier, and qr_token columns.")
    event_id = str(uuid.uuid4())
    root = DATA_ROOT / "events" / event_id
    (root / "assets").mkdir(parents=True, exist_ok=False)
    try:
        saved_names: set[str] = set()
        stored_bytes = 0
        for supplied_name, content in avatar_files:
            relative = safe_relative_path(supplied_name)
            if relative.suffix.lower() not in ALLOWED_AVATAR_EXTENSIONS:
                continue  # Folder pickers also include files such as .DS_Store.
            if relative.as_posix().casefold() in saved_names:
                raise BadgeError("Duplicate avatar filename in the import.")
            saved_names.add(relative.as_posix().casefold())
            destination = security.asset_path(root / "assets", relative.as_posix())
            try:
                security.normalize_image_file(relative.name, content, destination)
            except BadgeError as exc:
                raise BadgeError(f"Avatar '{relative.name}': {exc}") from exc
            stored_bytes += destination.stat().st_size
            if stored_bytes > security.MAX_REQUEST_BYTES:
                raise BadgeError("Normalized avatars exceed the 4 GiB import limit.")
        records: list[dict[str, Any]] = []
        for index, row in enumerate(reader, start=1):
            if index > security.MAX_RECORDS:
                raise BadgeError("Import at most 10,000 attendee rows at a time.")
            if None in row:
                raise BadgeError(f"CSV row {index + 1}: too many columns.")
            if not any((value or "").strip() for value in row.values()):
                continue
            records.append({
                "number": (row.get("number") or str(index)).strip(),
                "ticket_id": (row.get("ticket_id") or "").strip(),
                "display_name": (row.get("display_name") or "").strip(),
                "tier": (row.get("tier") or "").strip(),
                "qr_token": row.get("qr_token") or "",
            })
        templates = discover_templates()
        avatars = avatar_index(root / "assets")
        seen_ticket_ids: set[str] = set()
        for row_number, record in enumerate(records, start=2):
            if record["ticket_id"] in seen_ticket_ids:
                raise BadgeError(f"CSV row {row_number}: ticket_id '{record['ticket_id']}' is duplicated.")
            seen_ticket_ids.add(record["ticket_id"])
            prepare_record(record, row_number, templates, avatars)
        event = {"id": event_id, "created_at": now(), "records": records}
        write_json(root / "event.json", event)
        return event_summary(event, templates)
    except Exception:
        shutil.rmtree(root, ignore_errors=True)
        raise


def event_summary(event: dict[str, Any], templates: dict[str, dict[str, Any]] | None = None) -> dict[str, Any]:
    templates = discover_templates() if templates is None else templates
    event_dir = event_directory(event["id"])
    badges = []
    missing_avatars = []
    for record in event["records"]:
        badge = {key: record[key] for key in ("ticket_id", "display_name", "tier")}
        badges.append(badge)
        # Use the same current Tier / saved template fallback as render_badge().
        template = template_for_tier(record["tier"], templates) or templates.get(record.get("template_id", ""))
        if not template or not template["manifest"]["profile_picture"]:
            continue
        reference = str(record.get("avatar_asset", "")).strip()
        try:
            present = bool(reference) and avatar_path(event_dir, reference).is_file()
        except (BadgeError, OSError):
            present = False
        if not present:
            missing_avatars.append(badge)
    return {
        "id": event["id"],
        "created_at": event["created_at"],
        "record_count": len(event["records"]),
        "badges": badges,
        "missing_avatars": missing_avatars,
    }


def load_event(event_id: str) -> tuple[Path, dict[str, Any]]:
    root = event_directory(event_id)
    return root, read_json(root / "event.json")


# Font selection shared by editor text and final badges. Draft previews can supply
# unsaved font definitions; exports read the saved manifest. Missing styles fall back to regular.
def _font_asset(element: dict[str, Any], root: Path, fonts: list[dict[str, Any]] | None = None) -> Path | None:
    asset = element.get("font_asset")
    if element.get("font_family"):
        families = fonts if fonts is not None else read_json(root / "template.json").get("fonts", [])
        family = next((item for item in families if item.get("id") == element["font_family"]), None)
        if family:
            asset = family.get(element.get("font_style", "regular")) or family.get("regular")
    if asset:
        path = security.asset_path(root, asset)
        if path.is_file():
            return path
    return None


def _composite_transformed(canvas: Image.Image, layer: Image.Image, x: int, y: int, rotation: float = 0) -> None:
    if rotation:
        angle = math.radians(rotation)
        security.check_dimensions(math.ceil(abs(layer.width * math.cos(angle)) + abs(layer.height * math.sin(angle))) + 2,
                                  math.ceil(abs(layer.width * math.sin(angle)) + abs(layer.height * math.cos(angle))) + 2)
        transformed = layer.rotate(-rotation, resample=Image.Resampling.BICUBIC, expand=True)
        x -= (transformed.width - layer.width) // 2
        y -= (transformed.height - layer.height) // 2
        layer = transformed
    canvas.alpha_composite(layer, (round(x), round(y)))


# TYPOGRAPHY CORE: editor and final output use the same Unicode fallback and shaping.
def _text_layer(element: dict[str, Any], value: str, template_root: Path, fonts: list[dict[str, Any]] | None = None) -> Image.Image:
    try:
        return typography.render_text(element, value, _font_asset(element, template_root, fonts), DATA_ROOT / "google-fonts")
    except google_fonts.FontError as exc:
        raise BadgeError(str(exc)) from exc


def _draw_text(canvas: Image.Image, element: dict[str, Any], value: str, template_root: Path, origin: int) -> None:
    scratch = _text_layer(element, value, template_root)
    x = origin + mm_to_px(element["x_mm"], element["_ppi"])
    y = origin + mm_to_px(element["y_mm"], element["_ppi"])
    _composite_transformed(canvas, scratch, x, y, float(element.get("rotation_degrees", 0)))


# Edit layout sends draft settings here. Validate them and render only the text area;
# no layout is saved. The frontend positions and rotates this transparent PNG.
@security.image_task
def editor_text_preview(template_id: str, payload: dict[str, Any]) -> bytes:
    template = discover_templates().get(template_id)
    if not template:
        raise BadgeError("Template not found.")
    element = payload.get("element")
    value = payload.get("value", "")
    if not isinstance(element, dict) or element.get("type") != "text" or not isinstance(value, str):
        raise BadgeError("Choose a text layer and valid preview content.")
    if len(value) > 8192:
        raise BadgeError("Text preview content is too long.")
    element = {**element, "x_mm": 0, "y_mm": 0, "rotation_degrees": 0}
    for key, low, high in (("width_mm", .01, 2000), ("height_mm", .01, 2000),
                           ("font_size_pt", 1, 1000), ("min_font_size_pt", 1, 1000),
                           ("max_lines", 1, 1000), ("letter_spacing_pt", -1000, 1000),
                           ("line_spacing_pt", -1000, 1000)):
        if key in element:
            number = element[key]
            if not isinstance(number, (int, float)) or not math.isfinite(number) or not low <= number <= high:
                raise BadgeError(f"Invalid text setting: {key}.")
    manifest = {**template["manifest"], "elements": [element], "profile_picture": False,
                "fonts": payload.get("fonts", template["manifest"].get("fonts", []))}
    validate_template(manifest, template["root"])
    element["_ppi"] = int(manifest["print"]["ppi"])
    if mm_to_px(element["width_mm"], element["_ppi"]) * mm_to_px(element["height_mm"], element["_ppi"]) > 16_000_000:
        raise BadgeError("The text area is too large to preview.")
    try:
        image = _text_layer(element, value, template["root"], manifest["fonts"])
        output = io.BytesIO()
        image.save(output, "PNG")
        return output.getvalue()
    except (OSError, ValueError, TypeError) as exc:
        raise BadgeError("The text preview could not be rendered. Check its font and settings.") from exc


# Normalize masks once for both CSS previews and badge rendering. Browsers otherwise
# treat opaque grayscale PNGs as fully visible alpha masks, ignoring their black areas.
def avatar_mask_image(path: Path) -> Image.Image:
    source, _ = security.decode_image(path)
    with source:
        if "A" in source.getbands():
            return source.convert("RGBA").getchannel("A")
        return source.convert("L")


@security.image_task
def template_mask_preview(template_id: str, asset_path: str) -> bytes:
    path = template_asset_path(template_id, asset_path)
    try:
        mask = avatar_mask_image(path)
        image = Image.new("RGBA", mask.size, "white")
        image.putalpha(mask)
        output = io.BytesIO()
        image.save(output, "PNG")
        return output.getvalue()
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise BadgeError("The avatar mask could not be read as an image.") from exc


# Avatar layer: resize with cover/contain, apply crop offsets and an optional mask,
# then composite at the layer position. Missing or unreadable avatars are skipped.
def _draw_avatar(canvas: Image.Image, element: dict[str, Any], record: dict[str, Any], event_dir: Path, template_root: Path, origin: int, ppi: int) -> None:
    reference = str(record.get(element["field"], "")).strip()
    if not reference:
        return
    try:
        path = avatar_path(event_dir, reference)
    except BadgeError:
        return
    if not path.is_file():
        return
    try:
        decoded, _ = security.decode_image(path, max_pixels=security.MAX_IMAGE_PIXELS)
        source = decoded.convert("RGBA")
        decoded.close()
    except (UnidentifiedImageError, OSError):
        return
    width = mm_to_px(element["width_mm"], ppi)
    height = mm_to_px(element["height_mm"], ppi)
    if element.get("fit", "cover") == "contain":
        base_scale = min(width / source.width, height / source.height)
    else:
        base_scale = max(width / source.width, height / source.height)
    scale = base_scale * max(.01, float(element.get("scale", 1)))
    resized_size = (max(1, round(source.width * scale)), max(1, round(source.height * scale)))
    security.check_dimensions(*resized_size)
    resized = source.resize(resized_size, Image.Resampling.LANCZOS)
    portrait = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    left = round((width - resized.width) / 2 + float(element.get("offset_x_pct", 0)) / 100 * width)
    top = round((height - resized.height) / 2 + float(element.get("offset_y_pct", 0)) / 100 * height)
    portrait.alpha_composite(resized, (left, top))
    mask_asset = element.get("mask_asset")
    if mask_asset:
        mask = avatar_mask_image(security.asset_path(template_root, mask_asset))
        mask = mask.resize((width, height), Image.Resampling.LANCZOS)
        portrait.putalpha(ImageChops.multiply(portrait.getchannel("A"), mask))
    x = origin + mm_to_px(element["x_mm"], ppi)
    y = origin + mm_to_px(element["y_mm"], ppi)
    _composite_transformed(canvas, portrait, x, y, float(element.get("rotation_degrees", 0)))


def _draw_artwork(canvas: Image.Image, element: dict[str, Any], template_root: Path, origin: int, ppi: int) -> None:
    try:
        decoded, _ = security.decode_image(security.asset_path(template_root, element["asset"]))
        source = decoded.convert("RGBA")
        decoded.close()
    except (KeyError, FileNotFoundError):
        return
    width = mm_to_px(element["width_mm"], ppi)
    height = mm_to_px(element["height_mm"], ppi)
    source = source.resize((width, height), Image.Resampling.LANCZOS)
    # Artwork may extend past the canvas when freely positioned. Preserve signed offsets;
    # mm_to_px() clamps dimensions to at least one pixel and is unsuitable for coordinates.
    x = origin + round(float(element["x_mm"]) / 25.4 * ppi)
    y = origin + round(float(element["y_mm"]) / 25.4 * ppi)
    _composite_transformed(canvas, source, x, y, float(element.get("rotation_degrees", 0)))


def _draw_qr(canvas: Image.Image, element: dict[str, Any], value: str, origin: int, ppi: int) -> None:
    if not value:
        return
    image = barcode_image(value, element.get("symbology", "qr"), element.get("foreground_color", "#000000"), element.get("background_color", "#ffffff"), element.get("quiet_zone_modules", 4), str(element.get("error_correction", "M")))
    width = mm_to_px(element["width_mm"], ppi)
    height = mm_to_px(element["height_mm"], ppi)
    scale = min(width // image.width, height // image.height)
    if scale < 1:
        raise BadgeError("Barcode area is too small for this content. Enlarge the barcode or shorten its content.")
    image = image.resize((image.width * scale, image.height * scale), Image.Resampling.NEAREST).convert("RGBA")
    x = origin + mm_to_px(element["x_mm"], ppi) + (width - image.width) // 2
    y = origin + mm_to_px(element["y_mm"], ppi) + (height - image.height) // 2
    _composite_transformed(canvas, image, x, y, float(element.get("rotation_degrees", 0)))


# BADGE RENDERING CORE: build the RGB master used by full previews and every export.
# Resolve the current Tier override so events imported before a layout edit see that edit.
@security.image_task
def render_badge(event_dir: Path, record: dict[str, Any], templates: dict[str, dict[str, Any]] | None = None) -> tuple[Image.Image, dict[str, Any]]:
    templates = templates or discover_templates()
    template = template_for_tier(str(record.get("tier", "")), templates) or templates.get(record.get("template_id", ""))
    if not template:
        raise BadgeError("The record's template is unavailable.")
    manifest, root = template["manifest"], template["root"]
    validate_template(manifest, root)
    printing = manifest["print"]
    ppi = int(printing["ppi"])
    bleed = mm_to_px(printing["bleed_mm"], ppi)
    width = mm_to_px(printing["width_mm"], ppi) + bleed * 2
    height = mm_to_px(printing["height_mm"], ppi) + bleed * 2
    background = manifest.get("background", {})
    canvas = Image.new("RGBA", (width, height), background.get("color", "#ffffff"))
    if background.get("asset"):
        background_path = security.asset_path(root, background["asset"])
        if not background_path.is_file():
            raise BadgeError(f"Background asset is missing: {background['asset']}")
        decoded, _ = security.decode_image(background_path)
        try:
            if decoded.size != (width, height):
                resized = decoded.resize((width, height), Image.Resampling.LANCZOS)
                decoded.close()
                decoded = resized
            if decoded.mode == "RGB":
                canvas.paste(decoded, (0, 0))
            else:
                canvas.alpha_composite(decoded)
        finally:
            decoded.close()
    # List order is paint order: later layers appear above earlier ones. Coordinates
    # are relative to the trim area; the renderer adds the bleed offset.
    for raw_element in manifest["elements"]:
        element = dict(raw_element)
        if not element.get("visible", True):
            continue
        element["_ppi"] = ppi
        if element["type"] == "text":
            value = element["text"] if "text" in element else str(record.get(element["field"], ""))
            _draw_text(canvas, element, value, root, bleed)
        elif element["type"] == "masked-image":
            _draw_avatar(canvas, element, record, event_dir, root, bleed, ppi)
        elif element["type"] == "qr":
            _draw_qr(canvas, element, str(record.get(element["field"], "")), bleed, ppi)
        elif element["type"] == "artwork":
            _draw_artwork(canvas, element, root, bleed, ppi)
    result = canvas.convert("RGB")
    canvas.close()
    return result, printing


@security.image_task
def preview(event_id: str, ticket_id: str) -> bytes:
    event_dir, event = load_event(event_id)
    record = next((item for item in event["records"] if item["ticket_id"] == ticket_id), None)
    if not record:
        raise BadgeError("Ticket ID not found.")
    image, printing = render_badge(event_dir, record)
    buffer = io.BytesIO()
    ppi = int(printing["ppi"])
    image.save(buffer, "PNG", dpi=(ppi, ppi))
    return buffer.getvalue()


def _cmyk_image(image: Image.Image, profile_bytes: bytes) -> tuple[Image.Image, bytes]:
    security.check_size(profile_bytes, security.MAX_PROFILE_BYTES, "ICC profile")
    try:
        destination = ImageCms.ImageCmsProfile(io.BytesIO(profile_bytes))
        source = ImageCms.createProfile("sRGB")
        converted = ImageCms.profileToProfile(image.convert("RGB"), source, destination, outputMode="CMYK")
        return converted, destination.tobytes()
    except Exception as exc:
        raise BadgeError("The supplied ICC profile could not be used for RGB-to-CMYK conversion.") from exc


def artifact_stem(ticket_id: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", ticket_id).strip("._") or "badge"
    return cleaned


@security.image_task
def export_badge(event_id: str, ticket_id: str, kind: str, profile_bytes: bytes | None = None) -> tuple[bytes, str]:
    if kind not in ("jpg", "pdf"):
        raise BadgeError("Choose JPG or PDF export.")
    event_dir, event = load_event(event_id)
    record = next((item for item in event["records"] if item["ticket_id"] == ticket_id), None)
    if record is None:
        raise BadgeError("Ticket ID not found.")
    master, printing = render_badge(event_dir, record)
    converted, profile = _export_image(master, profile_bytes)
    jpeg = io.BytesIO()
    ppi = int(printing["ppi"])
    converted.save(jpeg, "JPEG", quality=95, subsampling=0, dpi=(ppi, ppi), icc_profile=profile)
    filename = f"{artifact_stem(ticket_id)}.{kind}"
    if kind == "jpg":
        return jpeg.getvalue(), filename
    width = mm_to_pt(float(printing["width_mm"]) + 2 * float(printing["bleed_mm"]))
    height = mm_to_pt(float(printing["height_mm"]) + 2 * float(printing["bleed_mm"]))
    pdf = io.BytesIO()
    document = pdf_canvas.Canvas(pdf, pagesize=(width, height), pageCompression=1)
    document.setTitle(f"Convention Badge {ticket_id}")
    document.drawImage(ImageReader(jpeg), 0, 0, width=width, height=height, mask=None)
    document.showPage()
    document.save()
    return pdf.getvalue(), filename


def export_cmyk_jpeg_zip(event_id: str, profile_filename: str, profile_bytes: bytes) -> Path:
    if not profile_filename.lower().endswith((".icc", ".icm")):
        raise BadgeError("Choose an ICC or ICM profile for CMYK export.")
    return _export_jpeg_zip(event_id, profile_bytes)


def export_rgb_jpeg_zip(event_id: str) -> Path:
    return _export_jpeg_zip(event_id)


# Color conversion boundary: render in RGB first, then optionally apply the printer ICC.
# Both single-badge and batch exports use this conversion.
def _export_image(image: Image.Image, profile_bytes: bytes | None) -> tuple[Image.Image, bytes]:
    if profile_bytes is not None:
        return _cmyk_image(image, profile_bytes)
    return image.convert("RGB"), ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()


# Batch ZIP: one JPEG and one single-page PDF per Ticket ID, with collision checks.
@security.image_task
def _export_jpeg_zip(event_id: str, profile_bytes: bytes | None = None,
                     progress: Callable[[int, int, str], None] | None = None) -> Path:
    color_mode = "cmyk" if profile_bytes is not None else "rgb"
    event_dir, event = load_event(event_id)
    if not event["records"]:
        raise BadgeError("There are no badges to export.")
    filenames: set[str] = set()
    for record in event["records"]:
        stem = artifact_stem(record["ticket_id"]).casefold()
        if stem in filenames:
            raise BadgeError(f"Ticket ID '{record['ticket_id']}' produces a duplicate filename. Use distinct filename-safe Ticket IDs.")
        filenames.add(stem)
    templates = discover_templates()
    export_dir = event_dir / "exports"
    export_dir.mkdir(exist_ok=True)
    filename = f"badge-export-{color_mode}-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:8]}.zip"
    output = export_dir / filename
    try:
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            if progress:
                progress(0, len(event["records"]), "rendering")
            for index, record in enumerate(event["records"], 1):
                master, printing = render_badge(event_dir, record, templates)
                ppi = int(printing["ppi"])
                master_dir = event_dir / "masters"
                master_dir.mkdir(exist_ok=True)
                master.save(master_dir / f"{artifact_stem(record['ticket_id'])}.png", "PNG", dpi=(ppi, ppi))
                converted, embedded_profile = _export_image(master, profile_bytes)
                blob = io.BytesIO()
                converted.save(blob, "JPEG", quality=95, subsampling=0, dpi=(ppi, ppi), icc_profile=embedded_profile)
                stem = artifact_stem(record['ticket_id'])
                archive.writestr(f"{color_mode}-jpeg/{stem}.jpg", blob.getvalue())
                page_width = mm_to_pt(float(printing["width_mm"]) + float(printing["bleed_mm"]) * 2)
                page_height = mm_to_pt(float(printing["height_mm"]) + float(printing["bleed_mm"]) * 2)
                pdf = io.BytesIO()
                document = pdf_canvas.Canvas(pdf, pagesize=(page_width, page_height), pageCompression=1)
                document.setTitle(f"Convention Badge {record['ticket_id']}")
                document.drawImage(ImageReader(blob), 0, 0, width=page_width, height=page_height, mask=None)
                document.showPage()
                document.save()
                archive.writestr(f"{color_mode}-pdf/{stem}.pdf", pdf.getvalue())
                master.close()
                converted.close()
                if progress:
                    progress(index, len(event["records"]), "rendering")
            if progress:
                progress(len(event["records"]), len(event["records"]), "packing")
    except Exception:
        output.unlink(missing_ok=True)
        raise
    return output


def export_cmyk_master_pdf(event_id: str, profile_filename: str, profile_bytes: bytes) -> Path:
    """Create one CMYK-image PDF page per master badge; this is not an imposition engine."""
    if not profile_filename.lower().endswith((".icc", ".icm")):
        raise BadgeError("Choose an ICC or ICM profile for CMYK export.")
    return _export_master_pdf(event_id, profile_bytes)


def export_rgb_master_pdf(event_id: str) -> Path:
    return _export_master_pdf(event_id)


# Master PDF: one badge per page at its physical size, including bleed; no sheet imposition.
@security.image_task
def _export_master_pdf(event_id: str, profile_bytes: bytes | None = None,
                       progress: Callable[[int, int, str], None] | None = None) -> Path:
    color_mode = "cmyk" if profile_bytes is not None else "rgb"
    event_dir, event = load_event(event_id)
    templates = discover_templates()
    export_dir = event_dir / "exports"
    export_dir.mkdir(exist_ok=True)
    output = export_dir / f"badge-masters-{color_mode}-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:8]}.pdf"
    document: pdf_canvas.Canvas | None = None
    try:
        if progress:
            progress(0, len(event["records"]), "rendering")
        for index, record in enumerate(event["records"], 1):
            master, printing = render_badge(event_dir, record, templates)
            converted, embedded_profile = _export_image(master, profile_bytes)
            jpeg = io.BytesIO()
            ppi = int(printing["ppi"])
            converted.save(jpeg, "JPEG", quality=95, subsampling=0, dpi=(ppi, ppi), icc_profile=embedded_profile)
            page_width = mm_to_pt(float(printing["width_mm"]) + float(printing["bleed_mm"]) * 2)
            page_height = mm_to_pt(float(printing["height_mm"]) + float(printing["bleed_mm"]) * 2)
            if document is None:
                document = pdf_canvas.Canvas(str(output), pagesize=(page_width, page_height), pageCompression=1)
                document.setTitle("Convention Badge Masters")
            else:
                document.setPageSize((page_width, page_height))
            document.drawImage(ImageReader(jpeg), 0, 0, width=page_width, height=page_height, mask=None)
            document.showPage()
            master.close()
            converted.close()
            if progress:
                progress(index, len(event["records"]), "rendering")
        if document is None:
            raise BadgeError("There are no badges to export.")
        if progress:
            progress(len(event["records"]), len(event["records"]), "packing")
        document.save()
        return output
    except Exception as exc:
        if document is not None and not isinstance(exc, ExportCancelled):
            try:
                document.save()
            except Exception:
                pass
        output.unlink(missing_ok=True)
        raise
